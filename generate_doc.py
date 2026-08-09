from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

doc = Document()

def set_rtl(paragraph):
    pPr = paragraph._p.get_or_add_pPr()
    bidi = OxmlElement('w:bidi')
    pPr.append(bidi)
    jc = OxmlElement('w:jc')
    jc.set(qn('w:val'), 'right')
    pPr.append(jc)
    for run in paragraph.runs:
        rPr = run._r.get_or_add_rPr()
        rtl = OxmlElement('w:rtl')
        rPr.append(rtl)

def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    set_rtl(p)
    return p

_numbered_counter = [0]

def add_bullet(doc, text):
    p = doc.add_paragraph()
    run = p.add_run(f'• {text}')
    pPr = p._p.get_or_add_pPr()
    bidi = OxmlElement('w:bidi')
    pPr.append(bidi)
    jc = OxmlElement('w:jc')
    jc.set(qn('w:val'), 'right')
    pPr.append(jc)
    ind = OxmlElement('w:ind')
    ind.set(qn('w:right'), '360')
    pPr.append(ind)
    rPr = run._r.get_or_add_rPr()
    rtl = OxmlElement('w:rtl')
    rPr.append(rtl)
    return p

def add_numbered(doc, text, reset=False):
    if reset:
        _numbered_counter[0] = 0
    _numbered_counter[0] += 1
    p = doc.add_paragraph()
    run = p.add_run(f'{_numbered_counter[0]}. {text}')
    pPr = p._p.get_or_add_pPr()
    bidi = OxmlElement('w:bidi')
    pPr.append(bidi)
    jc = OxmlElement('w:jc')
    jc.set(qn('w:val'), 'right')
    pPr.append(jc)
    ind = OxmlElement('w:ind')
    ind.set(qn('w:right'), '360')
    pPr.append(ind)
    rPr = run._r.get_or_add_rPr()
    rtl = OxmlElement('w:rtl')
    rPr.append(rtl)
    return p

def add_para(doc, text, bold=False):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    set_rtl(p)
    return p

def add_code(doc, text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = 'Courier New'
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x20, 0x20, 0x80)
    p.style = 'Normal'
    return p

# Title page
title = doc.add_heading('מערכת ניטור תנועה חכמה', 0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_rtl(title)

sub = doc.add_paragraph('Smart Motion Monitoring System')
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
sub.runs[0].font.size = Pt(14)

doc.add_paragraph('')
meta = doc.add_paragraph('פרויקט גמר – קורס: פיתוח תוכנה משובצת\nסטודנט: עמרי זמיר\nתאריך הגשה: יוני 2026')
meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_rtl(meta)

doc.add_page_break()

# 1. Introduction
add_heading(doc, '1. מבוא', 1)
add_para(doc, 'פרויקט זה מציג מערכת ניטור תנועה בזמן אמת המבוססת על IoT (Internet of Things). המערכת מדמה חיישן PIR (Passive Infrared Sensor) לגילוי תנועה, שולחת נתונים לענן, ומציגה אותם בלוח בקרה ויזואלי.')
add_para(doc, '')
add_para(doc, 'הבעיה: כיצד ניתן לנטר תנועה בחדר מרחוק, בזמן אמת, ללא נוכחות פיזית?')
add_para(doc, 'הפתרון: שימוש בחיישן תנועה (מדומה בפייתון) המחובר לפלטפורמת ענן ThingsBoard, המציגה את הנתונים בלוח בקרה אינטראקטיבי.')
add_para(doc, '')
add_para(doc, 'מטרות הפרויקט:', bold=True)
add_bullet(doc, 'סימולציה של חיישן PIR באמצעות פייתון')
add_bullet(doc, 'שליחת נתוני טלמטריה לענן בזמן אמת')
add_bullet(doc, 'הצגת הנתונים בלוח בקרה ויזואלי')

doc.add_paragraph('')

# 2. System Design
add_heading(doc, '2. עיצוב המערכת', 1)
add_para(doc, 'ארכיטקטורת המערכת:', bold=True)
add_para(doc, 'המערכת מורכבת משלושה רכיבים עיקריים המחוברים בזרימת נתונים רציפה:')
doc.add_paragraph('')
add_para(doc, '[ Python Simulator ]  →  [ ThingsBoard HTTP API ]  →  [ Dashboard ]')
doc.add_paragraph('')

add_para(doc, 'רכיבי המערכת:', bold=True)
add_bullet(doc, 'חיישן PIR מדומה – סקריפט פייתון המייצר אירועי תנועה אקראיים (הסתברות 40% לכל מרווח)')
add_bullet(doc, 'ThingsBoard Cloud – פלטפורמת IoT חינמית לאחסון ויזואליזציה של נתוני חיישנים')
add_bullet(doc, 'HTTP REST API – פרוטוקול התקשורת בין הסימולטור לענן')
add_bullet(doc, 'לוח בקרה – מציג גרף זמן-סדרה של אירועי תנועה ומונה סה"כ גילויים')

doc.add_paragraph('')
add_para(doc, 'זרימת נתונים:', bold=True)
add_para(doc, 'כל 3 שניות הסקריפט מחולל ערך תנועה (0 או 1), מצרף חותמת זמן ומונה, ושולח JSON ל-ThingsBoard דרך HTTP POST. ThingsBoard שומר את הנתונים ומעדכן את לוח הבקרה בזמן אמת.')

doc.add_paragraph('')

# 3. Implementation
add_heading(doc, '3. מימוש', 1)

add_heading(doc, '3.1 סקריפט הסימולטור – sensor_simulator.py', 2)
add_para(doc, 'הסקריפט כתוב בפייתון ומורכב משני פונקציות מרכזיות:')
doc.add_paragraph('')

add_para(doc, 'פונקציה 1 – simulate_pir_sensor():', bold=True)
add_para(doc, 'מדמה חיישן PIR על ידי הגרלה אקראית. מחזירה 1 (תנועה) ב-40% מהמקרים ו-0 (אין תנועה) ב-60% מהמקרים.')
add_code(doc, 'def simulate_pir_sensor():\n    return 1 if random.random() < MOTION_PROBABILITY else 0')

doc.add_paragraph('')
add_para(doc, 'פונקציה 2 – send_to_thingsboard(motion, total):', bold=True)
add_para(doc, 'שולחת את נתוני הטלמטריה ל-ThingsBoard דרך HTTP POST. הפייילואד מכיל: ערך תנועה, סה"כ גילויים, וחותמת זמן.')
add_code(doc, 'def send_to_thingsboard(motion, total):\n    payload = {\n        "motion": motion,\n        "total_detections": total,\n        "timestamp": datetime.now().strftime("%Y-%m-%d %H:%M:%S")\n    }\n    response = requests.post(TB_URL, json=payload, timeout=5)\n    return response.status_code == 200')

doc.add_paragraph('')
add_heading(doc, '3.2 ניהול אישורים – .env', 2)
add_para(doc, 'טוקן הגישה ל-ThingsBoard נשמר בקובץ .env נפרד ולא בקוד, למניעת חשיפה בטעות. הספרייה python-dotenv טוענת אותו בזמן ריצה:')
add_code(doc, 'TB_TOKEN=<your_device_token>')
add_code(doc, 'TB_TOKEN = os.getenv("TB_TOKEN")')

doc.add_paragraph('')
add_heading(doc, '3.3 הגדרת ThingsBoard', 2)
add_numbered(doc, 'יצירת חשבון חינמי ב-ThingsBoard Cloud', reset=True)
add_numbered(doc, 'יצירת מכשיר (Device) בשם "MotionSensor"')
add_numbered(doc, 'העתקת טוקן הגישה לקובץ .env')
add_numbered(doc, 'יצירת Dashboard עם שתי ווידג\'טים: גרף זמן-סדרה ומונה')

doc.add_paragraph('')

# 4. Testing
add_heading(doc, '4. בדיקות ותוצאות', 1)
add_para(doc, 'הרצת הסימולטור:', bold=True)
add_code(doc, 'pip install -r requirements.txt\npython sensor_simulator.py')
doc.add_paragraph('')
add_para(doc, 'פלט הטרמינל (דוגמה):', bold=True)
add_code(doc, '=== Motion Sensor Simulator Started ===\nSending data to ThingsBoard every 3 seconds...\n\n[14:32:01] No motion.\n  -> Sent to ThingsBoard OK\n[14:32:04] MOTION DETECTED! Total: 1\n  -> Sent to ThingsBoard OK\n[14:32:07] No motion.\n  -> Sent to ThingsBoard OK\n[14:32:10] MOTION DETECTED! Total: 2\n  -> Sent to ThingsBoard OK')
doc.add_paragraph('')
add_para(doc, 'תוצאות:', bold=True)
add_bullet(doc, 'הסימולטור שלח נתונים ל-ThingsBoard בהצלחה בכל מרווח של 3 שניות')
add_bullet(doc, 'לוח הבקרה הציג את אירועי התנועה בגרף זמן-סדרה בזמן אמת')
add_bullet(doc, 'מונה הגילויים התעדכן בהתאם לאירועי התנועה')
add_bullet(doc, 'אחוז גילוי בפועל היה קרוב ל-40% כצפוי מההגדרה')
doc.add_paragraph('')
add_para(doc, '[כאן יש להוסיף צילומי מסך של לוח הבקרה]', bold=False)

doc.add_paragraph('')

# 5. Conclusion
add_heading(doc, '5. סיכום', 1)
add_para(doc, 'הפרויקט הדגים בהצלחה מערכת IoT מקצה לקצה לניטור תנועה. הושגו כל מטרות הפרויקט:')
add_bullet(doc, 'סימולציה מציאותית של חיישן PIR בפייתון')
add_bullet(doc, 'שידור נתונים לענן בפרוטוקול HTTP REST')
add_bullet(doc, 'ויזואליזציה בזמן אמת בלוח בקרה מקצועי')

doc.add_paragraph('')
add_para(doc, 'מה למדתי:', bold=True)
add_bullet(doc, 'פרוטוקול HTTP ועבודה עם REST API')
add_bullet(doc, 'פלטפורמות IoT וניהול נתוני חיישנים בענן')
add_bullet(doc, 'ניהול אישורים בצורה מאובטחת')
add_bullet(doc, 'עיצוב ויזואליזציה של נתונים בזמן אמת')

doc.add_paragraph('')
add_para(doc, 'עבודה עתידית:', bold=True)
add_bullet(doc, 'חיבור לחיישן PIR אמיתי דרך ESP32 או Raspberry Pi')
add_bullet(doc, 'הוספת התראות במייל או SMS בעת גילוי תנועה')
add_bullet(doc, 'שמירת היסטוריה ואנליזה של דפוסי תנועה לאורך זמן')

output_path = r'C:\Users\omriz\OneDrive\שולחן העבודה\motion-monitor\Smart_Motion_Monitoring_System.docx'
doc.save(output_path)
print(f"Document saved to: {output_path}")
